from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/proposal';OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)

NAVY='10233F';BLUE='276EF1';LIGHT='EAF2FF';PALE='F4F7FB';GRID='D9D9D9';BLACK='000000';GRAY='4B5563';WHITE='FFFFFF'
def font(run,name='Aptos',size=10,color=BLACK,bold=False):
    run.font.name=name;run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name);run._element.rPr.rFonts.set(qn('w:hAnsi'),name);run.font.size=Pt(size);run.font.color.rgb=RGBColor.from_string(color);run.bold=bold
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd');tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def margins(cell,top=60,start=90,bottom=60,end=90):
    tcPr=cell._tc.get_or_add_tcPr();tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:tcMar=OxmlElement('w:tcMar');tcPr.append(tcMar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+tag))
        if node is None:node=OxmlElement('w:'+tag);tcMar.append(node)
        node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa')
def borders(table,color=GRID):
    tblPr=table._tbl.tblPr;b=tblPr.first_child_found_in('w:tblBorders')
    if b is None:b=OxmlElement('w:tblBorders');tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),color);b.append(el)
def set_cell(cell,text,bold=False,color=BLACK,size=8.7,align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text='';p=cell.paragraphs[0];p.alignment=align;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
    r=p.add_run(str(text));font(r,size=size,color=color,bold=bold);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;margins(cell)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;borders(t)
    for i,h in enumerate(headers):
        if widths:t.columns[i].width=Inches(widths[i])
        set_cell(t.rows[0].cells[i],h,True,WHITE,8.4,WD_ALIGN_PARAGRAPH.CENTER);shade(t.rows[0].cells[i],NAVY)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell(cells[i],val,False,BLACK,8.0,WD_ALIGN_PARAGRAPH.CENTER if i>0 and len(str(val))<22 else WD_ALIGN_PARAGRAPH.LEFT)
            if ri%2:shade(cells[i],PALE)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def add_p(text='',bold_start=None,space=5):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(space);p.paragraph_format.line_spacing=1.12
    if bold_start and text.startswith(bold_start):
        r=p.add_run(bold_start);font(r,bold=True);r=p.add_run(text[len(bold_start):]);font(r)
    else:font(p.add_run(text))
    return p
def bullets(items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(2);p.paragraph_format.left_indent=Inches(.22);font(p.add_run(x),size=9.5)
def heading(text,level=1):
    p=doc.add_heading(text,level=level);p.paragraph_format.keep_with_next=True;return p
def page():doc.add_page_break()
def image(path,width,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Inches(width))
    cp=doc.add_paragraph();cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.space_after=Pt(10);font(cp.add_run(caption),size=8,color=GRAY)

styles=doc.styles
styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10);styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
for name,size in [('Title',30),('Heading 1',20),('Heading 2',13),('Heading 3',11)]:
    st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=True;st._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display');st._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos Display');st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(6)
styles['Title'].paragraph_format.space_before=Pt(0);styles['Title'].paragraph_format.space_after=Pt(10)

# Cover
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ROOT/'public/brand/logo.png'),width=Inches(2.45))
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('AIDA Data Rescue Hackathon Proposal')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('Ajman University'),size=16,bold=True);p.paragraph_format.space_after=Pt(4)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('From Messy Data to Machine Learning'),size=12,color=GRAY)
doc.add_paragraph().paragraph_format.space_after=Pt(4)
image(ROOT/'output/mockups/welcome-concept.png',6.45,'Concept image for the proposed participant experience')
table(['Decision requested','Proposed event'],[['Approval','University lab, 45 students, sponsor outreach'],['Format','15 teams of three using one computer per team'],['Duration','90 minutes, including 60 minutes of competition'],['Funding ceiling','AED 4,000; sponsorship preferred'],['Preparation','Four weeks after approval']],widths=[2.0,4.8])
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('Prepared by the AI and Data Analytics Club'),size=9,color=GRAY)

page();heading('Executive Summary')
add_p('The AI and Data Analytics Club requests approval to run AIDA Data Rescue, a 90-minute educational hackathon for 45 Ajman University students working in 15 teams. Participants will clean synthetic campus-equipment data, construct labels from maintenance records, and train a small neural network through three timed stages. The competition is designed for students with basic Python and pandas experience; starter code keeps the focus on data quality and model improvement.')
add_p('We request a computer lab with 15 team computers and two spares, internet access, a projector, a faculty supervisor, two judges, one technical operator, and three student volunteers. We also request permission to approach sponsors through university-approved channels. The proposed spending ceiling is AED 4,000. No funding, sponsorship, venue, or infrastructure is currently assumed to be confirmed.')
heading('Approval Requested',2)
bullets(['Approve the event concept, participant capacity, format, and four-week preparation period.','Reserve a suitable university computer lab and technical support for a rehearsal and the event.','Authorize institutional sponsor outreach led by a faculty adviser, Student Life, or the Ajman University Innovation Center.','Approve a spending ceiling of AED 4,000, subject to university purchasing and sponsorship rules.'])
heading('Expected Learning Outcomes',2)
bullets(['Recognize missing values, malformed numeric data, duplicates, impossible readings, and identifier inconsistencies.','Create reliable labels by normalizing and joining a separate outcome log.','Prevent data leakage by fitting preprocessing only on the training partition.','Use validation accuracy and binary cross-entropy loss to compare neural-network experiments.','Explain technical decisions and work responsibly with limited AI assistance.'])
heading('Event Participants and Staff',2)
table(['Role','Count','Main responsibility'],[['Students','45','Compete in 15 teams of three'],['Faculty supervisor','1','Academic oversight and final decisions'],['Judges','2','Review evidence and approve scores'],['Technical operator','1','Manage clock, platform, incidents, and judging'],['Student volunteers','3','Registration, room support, and participant questions']],widths=[1.65,.75,4.4])

page();heading('Competition Format')
add_p('The official competition lasts 60 minutes. Three stages run for 20 minutes each, followed by 30 minutes for final scoring, provisional results, awards, and photographs. Registration, login checks, the guide, and environment checks happen before the official clock starts.')
table(['Time','Activity','AI access'],[['00:00 to 20:00','Stage 1 Rescue the Data','Unlocks at 05:00'],['20:00 to 40:00','Stage 2 Build the Labels','Unlocks at 25:00'],['40:00 to 60:00','Stage 3 Train and Improve','Unlocks at 45:00'],['60:00 to 70:00','Final scoring and verification','Closed'],['70:00 to 80:00','Publish provisional results','Closed'],['80:00 to 90:00','Awards and photographs','Closed']],widths=[1.35,3.8,1.65])
heading('Stage 1 Rescue the Data',2)
add_p('Students normalize identifiers and numerical values, remove duplicate copies, convert impossible readings to missing values, and document an imputation strategy. Learned imputation values must come from the fixed training partition.')
table(['Criterion','Points'],[['Schema, parsing, units, and ranges','10'],['Deduplication and identifier preservation','10'],['Missing-value handling and data-quality report','10']],widths=[5.6,1.2])
heading('Stage 2 Build the Labels',2)
add_p('Students clean and join a separate maintenance-outcome log. Corrective repair confirmed maps to 1; inspection complete with no corrective repair maps to 0. Unmatched records must be reported. Outcome text, record IDs, and target values cannot be model features.')
table(['Criterion','Points'],[['Correct labels matched to records','15'],['Fixed split and target-leakage prevention','10'],['Usable numerical input and consistent preprocessing','5']],widths=[5.6,1.2])
heading('Stage 3 Train and Improve',2)
add_p('Students begin with a 12 input, 16 hidden, 8 hidden, and 1 output PyTorch network. They may tune the learning rate, batch size, dropout, weight decay, hidden widths, and epoch count up to 50. Each training run is limited to 120 seconds on identical CPU resources.')
table(['Criterion','Points'],[['Private-test accuracy','20'],['Private-test binary cross-entropy loss','10'],['Reproducible model and preprocessing','5'],['Explanation of two experiments','5']],widths=[5.6,1.2])

page();heading('Dataset Design')
add_p('Campus Equipment Rescue is a fictional dataset created for education. It contains no student information, personal data, or confidential university records. The task is to predict whether a piece of equipment requires corrective maintenance within seven days.')
table(['Component','Specification'],[['Unique records','10,000'],['Participant development file','8,200 rows: 8,000 unique records plus 200 duplicate copies'],['Training and validation','6,000 training and 2,000 validation records'],['Private test set','2,000 records held by organizers'],['Input variables','12 numeric equipment readings'],['Target','needs_maintenance, created by students'],['Outcome log','8,000 development records'],['Observed positive rate','49.05 percent']],widths=[2.4,4.4])
heading('Input Variables',2)
table(['Variable','Meaning'],[[a,b] for a,b in [('temperature_c','Equipment temperature'),('ambient_temperature_c','Ambient temperature'),('vibration_mm_s','Vibration velocity'),('current_a','Electrical current'),('voltage_v','Electrical voltage'),('humidity_pct','Relative humidity'),('pressure_bar','Operating pressure'),('rotation_rpm','Rotation speed'),('daily_runtime_hours','Daily operating time'),('equipment_age_months','Equipment age'),('days_since_service','Time since service'),('load_pct','Operating load')]],widths=[2.65,4.15])
heading('Controlled Data Problems',2)
bullets(['Approximately 5 percent of feature cells are missing.','Some numeric values contain units or extra spaces.','A small number of readings exceed published physical ranges.','Two hundred duplicate copies are mixed into the development file.','Record identifiers and maintenance statuses use inconsistent spacing and capitalization.'])
add_p('The organizers retain a clean reference dataset, private test labels, and a corruption manifest. Participant downloads exclude those files.')

page();heading('Scoring and Fairness')
add_p('The judges compute scores from frozen submissions and actual predictions. Participants cannot score their own private-test results. A fresh isolated run must reproduce the submitted artifacts before the model receives reproducibility points.')
heading('Metric Formula',2)
add_p('Accuracy points = 20 x clamp((accuracy - 0.50) / 0.40, 0, 1). Loss points = 10 x clamp((ln(2) - loss) / (ln(2) - 0.30), 0, 1). Accuracy is expressed as a fraction. The clamp function limits each normalized value to the interval from zero to one.')
add_p('The reference implementation completed in approximately 4.4 seconds on a local CPU. It reached 88.8 percent validation accuracy, validation loss of 0.273, private-test accuracy of 87.7 percent, and private-test loss of 0.281. These results support a practice target of at least 85 percent validation accuracy and validation loss no greater than 0.40.')
heading('Tie Resolution',2)
add_p('Ties are resolved by private-test accuracy, then private-test loss, then the Stage 1 score. If teams remain tied, they share the placing and the related prize pool is divided equally.')
heading('Submission Rules',2)
bullets(['The server clock controls the shared event deadline and opens each scheduled stage for every team.','A valid Stage 1 or Stage 2 submission immediately opens the next stage for that team.','Teams may replace a submission only while that stage remains their active stage; the latest accepted receipt counts.','Autosave preserves work but does not count as submission.','At the deadline, the last explicitly submitted version is frozen.','A missed stage earns zero for that stage but does not block later work when the schedule reaches it.','Previous-stage code remains visible as a read-only snapshot.'])
heading('Manual Review',2)
add_p('Automated checks verify row coverage, schema, ranges, labels, preprocessing metadata, model execution, and private-test metrics. Faculty judges review the imputation rationale and the explanation of two experiments. No leaderboard is published until the judges approve the evidence and the technical operator records a reason for every score.')

page();heading('Limited AI Assistance')
add_p('Each team receives one built-in AI prompt in each stage, for a maximum of three prompts per team and 45 successful responses across the event. The assistant supports learning by explaining a concept, interpreting a short error, or suggesting an experiment. It cannot access the dataset, run code, edit files, or complete the full pipeline.')
table(['Rule','Limit'],[['Unlock time','Five minutes after each stage begins'],['Prompt size','80 whitespace-separated words and 1,000 characters'],['Response size','160 displayed words'],['Illustrative code','No more than five lines'],['Follow-up conversation','Not available'],['Unused prompts','Do not carry forward'],['Model settings','Same organizer-selected model for all teams']],widths=[2.25,4.55])
heading('Control Measures',2)
bullets(['The server atomically claims one prompt for each team and stage.','Duplicate clicks return the same claimed request state.','A provider failure allows the same prompt to be retried; it does not create a second prompt.','The system instruction restricts the assistant to conceptual help and a short response.','Only the built-in assistant is permitted on supervised university computers.','If AI fails for every team, the organizer pauses the shared event clock and logs the incident.'])
add_p('Prompt and response limits reduce the chance that the assistant completes the work, but they cannot guarantee this. The scoring rubric therefore rewards reproducibility and a short explanation of the team’s own experiments.')

page();heading('Participant Platform')
add_p('The implemented desktop-first website includes team login, dataset preview and downloads, a guide sheet, a Python editor, console output, training metrics, a server-synchronized timer, AI prompt controls, autosave status, submission receipts, prior-stage snapshots, and organizer controls.')
image(ROOT/'../assets/aida-landing-screenshot.png',6.65,'Implemented local website landing page')
image(ROOT/'../assets/aida-workspace-screenshot.png',6.65,'Implemented participant workspace')
page();heading('Platform Architecture')
heading('Visual Design',2)
add_p('The interface uses the supplied AIDA logo, deep navy surfaces, electric blue borders, cyan highlights, connected-node patterns, and the supplied mascot without recoloring it. Motion is disabled when the operating system requests reduced motion. The one-minute warning does not interrupt typing.')
table(['Subsystem','Implementation'],[['Web application','Next.js and TypeScript on Vercel'],['Code editor','Monaco Editor, packaged locally'],['Identity and durable state','Supabase Auth and Postgres'],['Submission artifacts','Private Supabase Storage bucket'],['Python execution','One isolated Vercel Sandbox per active team'],['Machine learning','Pinned CPU PyTorch environment'],['AI help','OpenAI Responses API, server-side only'],['Judging','Fresh isolated run followed by trusted inference and scoring']],widths=[2.2,4.6])
heading('Security and Integrity',2)
bullets(['Participant code runs in a remote microVM with outbound network access denied.','Application secrets and private test labels are never mounted into participant environments.','Team requests are authenticated on the server; one team cannot read another team’s work.','Supabase row-level security is enabled and direct browser access to event state and artifacts is revoked.','The OpenAI API key, Supabase service credential, and Vercel credential remain server-side.','The judging system supplies private test features to trusted inference only; scoring labels remain outside the execution sandbox.','A 120-second command timeout, two CPU threads, file allowlists, and 10 MB artifact limits constrain runs.','Audit records capture event-control actions and submission receipts.'])
heading('Hosting Requirements',2)
add_p('Vercel’s published Hobby allowance permits 10 concurrent sandboxes, which is below the 15-team requirement. The event therefore needs a paid Vercel allocation or approved university infrastructure. Supabase currently lists a free plan for rehearsal and a Pro plan starting at USD 25 per month. Final procurement must verify current prices and university rules before purchase.')
heading('Deployment Status',2)
add_p('The application, dataset generator, reference model, scoring scripts, and deployment configuration are complete locally. A cloud deployment is not active because the connected Supabase account contains no project and no Vercel project credentials were supplied. The project fails closed when those services are absent; participant code is never executed on the organizer’s computer.')

page();heading('Budget Request')
add_p('The figures below are planning allowances rather than supplier quotations. No expenditure is currently approved. University facilities, equipment, staff time, and platform development are requested in kind.')
table(['Item','Calculation','AED'],[['First-place prize','AED 400 x 3 students','1,200'],['Second-place prize','AED 250 x 3 students','750'],['Third-place prize','AED 150 x 3 students','450'],['Refreshments','55 people x AED 12','660'],['Trophies or medals','Allowance','150'],['Printing and signs','Allowance','100'],['Hosting and execution','Event and rehearsal allowance','200'],['LLM API usage','Event and testing allowance','50'],['Subtotal','','3,560'],['Contingency','10 percent','356'],['Estimated total','','3,916'],['Requested ceiling','Rounded','4,000']],widths=[2.5,3.25,1.05])
heading('Lower Cost Option',2)
add_p('A technical-only version can target AED 300 by using university facilities, digital certificates, no catering, and donated prizes. A zero-cash event requires confirmed in-kind compute, AI access, and prizes before registration opens.')
heading('Prize Conditions',2)
add_p('Monetary prizes should not be advertised until funding is secured. Product or voucher prizes should provide equal value to all three students on a winning team. Shared placings divide the relevant prize pool equally.')
heading('In Kind Requests',2)
bullets(['Computer lab with 15 team computers and two spares.','Stable internet connection and a projector.','Technical rehearsal access before the event.','Faculty supervisor, two judges, one technical operator, and three volunteers.','Institutional sponsor introductions and purchasing support.'])

page();heading('Sponsor Approach')
add_p('Sponsor outreach should start with university relationships. A faculty adviser, Student Life, or the Ajman University Innovation Center should make or approve introductions. The club should present the event as a pilot and avoid offering a long-term official sponsor title without university approval.')
table(['Organization','Proposed request','Contact route'],[['Ajman University Innovation Center','Introductions, mentors, judges, or infrastructure','auic@ajman.ac.ae; +971 6 705 6481'],['Ajman Bank','Approximately AED 2,400 for prizes','University partnership owner; then info@ajmanbank.ae'],['Mbank','Prizes, guest judge, or educational vouchers','University office managing the existing AU-Mbank relationship'],['AUIC startups and suppliers','Refreshments, printing, or technology gifts','Introduction through AUIC or Student Life']],widths=[1.75,2.8,2.25])
heading('Suggested Packages',2)
table(['Package','Request','Recognition subject to approval'],[['Lead sponsor','AED 2,400','Approved logo placement, acknowledgement, and award invitation'],['Supporting sponsor','AED 1,000','Acknowledgement and event invitation'],['Technology partner','AED 300 to 600 or in kind','Acknowledgement for compute, API credits, or vouchers']],widths=[1.6,1.7,3.5])
heading('Draft Sponsor Message',2)
add_p('Subject: Sponsorship request for AIDA Data Rescue at Ajman University')
add_p('Dear [Name],')
add_p('The AI and Data Analytics Club at Ajman University is seeking support for AIDA Data Rescue, a 90-minute educational hackathon for 45 students. Fifteen teams will clean a deliberately messy dataset, construct labels from evidence, and train a neural network in a supervised university lab.')
add_p('We invite [Organization] to support student prizes, refreshments, technology costs, or a guest judge. The proposed lead contribution is AED 2,400, which funds all three winning teams. Any recognition, logo placement, or public announcement will follow Ajman University approval.')
add_p('We would welcome a short meeting to share the scoring rubric, technical controls, schedule, and expected student outcomes. Thank you for considering this request.')
add_p('Sincerely,\nAI and Data Analytics Club\nAjman University')

page();heading('Preparation Schedule')
table(['Week','Work','Completion evidence'],[['1','Obtain approval; confirm venue and faculty roles; freeze dataset and reference solution','Signed approval path, room request, validated dataset'],['2','Configure cloud services; test login, timer, submissions, and AI quotas','Deployment checklist and service test results'],['3','Add final branding; prepare volunteers and judging workflow; contact sponsors','Rehearsal script, judge guide, sponsor status'],['4','Run a 15-team rehearsal; calibrate targets; freeze rules and funding','Signed readiness checklist and incident plan']],widths=[.7,3.75,2.3])
heading('Readiness Checklist',2)
bullets(['Fifteen teams can execute code concurrently with identical resources.','Beginner pilot teams can complete each stage within 20 minutes.','The baseline model trains within two minutes.','Cleaning improves data quality without target leakage.','Judges reproduce submissions and private-test scores.','Teams cannot read private labels or another team’s files.','AI unlocks and quotas survive duplicate requests, refreshes, and provider failures.','Deadlines reject late submissions and preserve accepted versions.','All submissions are judged within the first ten minutes after competition.','Animations do not interfere with editing or accessibility.','Spending alerts, environment shutdown, and the outage procedure have been rehearsed.'])
heading('Acceptance Results Completed Locally',2)
table(['Check','Result'],[['TypeScript and production build','Passed'],['Timing, prompt, scoring, and progression unit tests','8 of 8 passed'],['Authentication, team isolation, progression, and data-gating API checks','12 of 12 passed'],['Dependency security audit','0 known vulnerabilities after pinning'],['Reference model private-test accuracy','87.7 percent'],['Reference model private-test loss','0.281'],['Cloud concurrency rehearsal','Pending paid Vercel and Supabase configuration']],widths=[4.65,2.15])

page();heading('Risks and Controls')
table(['Risk','Control'],[['Cloud service interruption','Organizer pauses the shared clock, records the incident, and applies the same adjustment to every team.'],['Late or lost work','Five-second autosave, browser recovery copy, durable remote revisions, and explicit submission receipts.'],['AI assistant completes too much work','Short response, five-line code maximum, no tools, no dataset access, one prompt, and explanation points.'],['Data leakage','Fixed split, training-only preprocessing checks, private labels outside sandboxes, and trusted inference.'],['Unequal computing resources','One remote two-vCPU sandbox per team with the same image and timeout.'],['Sponsor funding is delayed','Do not advertise cash prizes; use the AED 300 technical-only option or postpone registration.'],['Student code is unsafe','Isolated microVMs, denied outbound network, no application secrets, timeouts, and artifact limits.'],['Scores are disputed','Frozen receipts, rerunnable code, audit records, published formulas, faculty review, and a ten-minute question period.']],widths=[2.05,4.75])
heading('Decision and Signatures')
add_p('Approval permits the club to begin the four-week preparation schedule and institutionally approved sponsor outreach. Event publication and prize advertising remain conditional on confirmed venue, technical capacity, judging staff, and funding.')
table(['Decision','Name and signature','Date'],[['Approved / Not approved','',''],['Dean or authorized representative','',''],['Faculty supervisor','',''],['Student Life representative','',''],['Technical readiness confirmation','','']],widths=[2.0,3.5,1.3])
heading('Administrative Details',2)
table(['Field','To be completed'],[['Event date',''],['College',''],['Venue',''],['Registration opening',''],['Registration closing',''],['Faculty supervisor',''],['Technical operator','']],widths=[2.6,4.2])

page();heading('Visual Concepts')
add_p('These generated images show the intended visual direction. They are concepts rather than screenshots of the working application. The preceding platform section contains screenshots of the implemented local website.')
for name,cap in [('workspace-concept.png','Concept for the participant workspace'),('guide-concept.png','Concept for the guide and time warning'),('finale-concept.png','Concept for the transition and results experience')]:image(ROOT/'output/mockups'/name,6.4,cap)

page();heading('Sources')
sources=[('Ajman University Innovation Center','https://auic.ajman.ac.ae/','Contact information and partnership role.'),('Ajman University and Ajman Bank','https://www.ajman.ac.ae/en/mou/ajman-bank','Existing university relationship.'),('Ajman Bank customer service charter','https://www.ajmanbank.ae/site/customer-service-charter.html','General inquiry contact route.'),('Ajman University and Mbank','https://www.ajman.ac.ae/en/news/2024/mbank-develops-strategic-partnership-with-ajman-university-to-offer-academic-and-professional-development-programs-to-university-students','Existing partnership concerning student development.'),('Vercel Sandbox pricing and quotas','https://vercel.com/docs/sandbox/pricing','Current concurrency, duration, and usage information; verify before purchase.'),('Supabase pricing','https://supabase.com/pricing','Current plan information; verify before purchase.'),('OpenAI Responses API','https://developers.openai.com/api/reference/resources/responses/methods/create','Server-side AI response configuration.')]
for title,url,note in sources:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7);r=p.add_run(title);font(r,bold=True);r=p.add_run('\n'+url+'\n'+note);font(r,size=8.5,color=GRAY)
add_p('Pricing and quotas can change. The organizing team must verify current provider terms, university procurement rules, sponsor approval, and data-protection requirements before deployment or purchase.',bold_start='Pricing and quotas can change.')

# Footer and page numbers
for section in doc.sections:
    footer=section.footer;p=footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('AIDA Data Rescue Hackathon Proposal   |   Ajman University   |   '),size=8,color=GRAY)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)

doc.core_properties.title='AIDA Data Rescue Hackathon Proposal'
doc.core_properties.subject='Approval, budget, sponsorship, technical design, and event rules'
doc.core_properties.author='AI and Data Analytics Club, Ajman University'
doc.save(OUT/'AIDA_Data_Rescue_Proposal.docx')
print(OUT/'AIDA_Data_Rescue_Proposal.docx')
